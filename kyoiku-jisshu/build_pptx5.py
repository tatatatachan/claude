import build_pptx4 as m4
m3=m4.m3; b=m3.b
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree
INK,SUB,TEAL,TEAL2,TEALD,ORG,ORGD=m3.INK,m3.SUB,m3.TEAL,m3.TEAL2,m3.TEALD,m3.ORG,m3.ORGD
card,num,note=m3.card,m3.circle_num,m3.note
n0=len(b.prs.slides)

# ---------- P24: 他社との比較 ----------
sl=m3.new_slide('採用戦略テーマ','他社と並べて見える、エフビー介護サービスの実習の魅力','P24：採用戦略テーマ（他社比較）。出典：1023採用定例資料。')
cols=['','エフビー介護サービス','敬老園','医療法人・緑風会','アズパートナー']
rows=[('ビジョン',['「共生」をテーマにした、人々が助け合う地域創生（独居社会の反対をつくる／エフビータウン構想・コミュニティタワー構想）','―','―','暮らしを、共に。']),
 ('事業フェーズ',['成長期〜拡大期（2022年 スタンダード市場上場）','安定期〜成熟期','地域密着・安定運営期','成長期〜拡大期（2024年 スタンダード市場上場）']),
 ('事業・特徴',['有料老人ホーム、グループホーム、デイ、訪問介護・看護、小規模多機能、福祉用具ほか。利用者の年齢にあわせて使える介護サービス＋小規模多機能','特養、ケアハウス、デイ、ショートステイ、居宅介護支援、地域包括支援センター。地域密着の総合福祉','病院運営、老健、デイケア、グループホーム、訪問看護ほか。病院母体の医療＋介護連携','有料老人ホーム「アズハイム」シリーズ（首都圏中心）。首都圏型・都市型の有料老人ホーム特化']),
 ('人・メンバー',['関わる人たちが、みんな優しく・温かい','「感謝の心を基にした関係」を築き、真心をもってサービスを提供できる人材','「利用者はもちろん、職員にとっても安心できる居場所」を重視','近年は毎年170名以上の新卒社員を受け入れ'])]
rh=[96,72,140,96]
gs=b.sh.add_table(5,5,b.px(64),b.px(176),b.px(1602),b.px(40+sum(rh))); tb=gs.table; gs.name='他社比較表'
tblPr=tb._tbl.tblPr
for a in ('firstRow','bandRow'): tblPr.set(a,'0')
sid=tblPr.find(qn('a:tableStyleId'))
if sid is not None: sid.text='{2D5ABB26-0587-4C30-8999-92F81FD0307C}'   # No Style, No Grid
widths=[150,420,344,344,344]
for i,w in enumerate(widths): tb.columns[i].width=b.px(w)
tb.rows[0].height=b.px(40)
for i,h in enumerate(rh): tb.rows[i+1].height=b.px(h)
def border(cell,color='DCE6E1',w=9525,sides=('B',)):
    tcPr=cell._tc.get_or_add_tcPr()
    for sd in sides:
        ln=etree.SubElement(tcPr,qn('a:ln'+sd),w=str(w),cap='flat',cmpd='sng',algn='ctr')
        sf=etree.SubElement(ln,qn('a:solidFill')); etree.SubElement(sf,qn('a:srgbClr'),val=color)
def setcell(c,txt,size,color,bold=False,fillc=None,align=PP_ALIGN.LEFT):
    c.margin_left=c.margin_right=b.px(14); c.margin_top=c.margin_bottom=b.px(8)
    c.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf=c.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]; p.alignment=align; p.line_spacing=1.2
    r=p.add_run(); r.text=txt; f=r.font; f.size=m3.b.FS(size); f.bold=bold; f.name=b.FONT; f.color.rgb=b.rgb(color)
    rPr=r._r.get_or_add_rPr()
    for tag in ('a:ea','a:cs'): etree.SubElement(rPr,qn(tag)).set('typeface',b.FONT)
    if fillc: c.fill.solid(); c.fill.fore_color.rgb=b.rgb(fillc)
    else: c.fill.background()
for j,h in enumerate(cols):
    c=tb.cell(0,j)
    setcell(c,h,19,'FFFFFF' if j==1 else INK,True,ORG if j==1 else 'EEF4F1',PP_ALIGN.CENTER)
for i,(lab,vals) in enumerate(rows,1):
    setcell(tb.cell(i,0),lab,19,TEALD,True,None)
    for j,v in enumerate(vals,1):
        c=tb.cell(i,j); setcell(c,v,17 if i==3 else 18,INK if j==1 else '55636B',False,'FFF4E8' if j==1 else None)
    for j in range(5): border(tb.cell(i,j))
top=176+40+sum(rh)+34
b.text(64,top,1500,40,'実習こそ、県内の他社には伝えられない魅力を、プログラムに落とし込む',27,INK,True,name='魅力見出し')
pts=[('成長産業の上場企業','その根本となる理念がある'),('ニーズに合わせたサービス展開','ご利用者さんに合わせた、実習経験の幅'),('ダイバーシティな環境','女性・技能実習生の採用など（例：県内の女性管理職の割合No.1）')]
for i,(h,d) in enumerate(pts):
    x=64+i*544; card(x,top+54,514,124)
    num(x+26,top+54+34,52,i+1,ORG if i==2 else TEAL2,26)
    b.text(x+96,top+54+22,396,44,h,22,TEALD,True,ls=1.1,name=f'魅力_{i+1}')
    b.text(x+96,top+54+64,396,50,d,17,INK,ls=1.25,name=f'魅力補足_{i+1}')
note('※ 他社の記載は公開情報の要約です。県内の女性管理職割合などの実績は、根拠資料の確認が必要です。')
m3.post(sl); m3.finish(sl)

# ---------- P27: 受け入れ体制 ----------
sl=m3.new_slide('実習の受け入れ体制','誰が担当しても伝わる、再現性の高い受け入れの仕組み','P27：実習の受け入れ体制（運用・仕組み）。')
p=b.rect(64,176,520,686,TEAL,MSO_SHAPE.ROUNDED_RECTANGLE,0.04); p.name='方針'
b.text(104,208,440,30,'実習コンテンツ（体制）',20,'EAF8F4',True,spacing=2,name='方針ラベル')
b.text(104,258,440,200,['誰が担当しても、','温かさとメッセージが','伝わる。'],38,'FFFFFF',True,ls=1.3,name='方針本文')
b.text(104,500,440,150,['「再現性の高い仕組み」を','つくります。','学生向けの学習シートなどを用意し、','プログラムを標準化します。'],21,'FFFFFF',ls=1.5,name='方針補足')
items=[('メンター体制','現場スタッフを「実習メンター」として配置。学生1〜2名につき1名が、親身に伴走します。',TEAL2),
 ('標準動画（ツール）','「ショート導入動画」などを活用。誰が説明しても、エフビーの温かさや社風がストレートに伝わります。',TEAL2),
 ('実習生の声の資産化','リアルな体験談・本音の感想をWeb上に蓄積・公開。次年度の学生や学校（教授・就職課）への開拓ツールに。',ORG),
 ('年間PDCAの循環','定期的に受け入れと振り返りを回し、学校や地域とのつながりを広げていきます。',ORG)]
for i,(h,d,c) in enumerate(items):
    x=620+(i%2)*532; y=176+(i//2)*350
    card(x,y,514,336)
    num(x+30,y+30,56,i+1,c,28)
    b.text(x+108,y+30,380,56,h,27,TEALD if c==TEAL2 else ORGD,True,anchor=MSO_ANCHOR.MIDDLE,name=f'体制見出し_{i+1}')
    b.text(x+36,y+120,444,190,d,25,INK,ls=1.5,name=f'体制説明_{i+1}')
note('出典：1023採用定例資料（実習の受け入れ体制）。')
m3.post(sl); m3.finish(sl)

# ---------- reorder ----------
lst=b.prs.slides._sldIdLst; ids=list(lst)
a,c=ids[n0],ids[n0+1]; lst.remove(a); lst.remove(c)
lst.insert(6,a); lst.insert(9,c)

# ---------- refine pass ----------
N=len(b.prs.slides)
for k,sld in enumerate(b.prs.slides,1):
    for shp in sld.shapes:
        if shp.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE:
            try: white=shp.fill.type==1 and str(shp.fill.fore_color.rgb)=='FFFFFF'
            except Exception: white=False
            if white and shp.auto_shape_type==MSO_SHAPE.ROUNDED_RECTANGLE:
                shp.adjustments[0]=min(0.5,20.0*b.E/min(shp.width,shp.height))
                shp.line.color.rgb=RGBColor.from_string('E1EAE6'); shp.line.width=Pt(1)
    tb_=sld.shapes.add_textbox(b.px(1546),b.px(894),b.px(120),b.px(26)); tf=tb_.text_frame
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT
    r=p.add_run(); r.text=f'{k} / {N}'; r.font.size=b.FS(15); r.font.name=b.FONT; r.font.color.rgb=RGBColor.from_string('8A9A96')
    tb_.name='ページ番号'
b.prs.save('共育実習_図解.pptx'); print(N)
