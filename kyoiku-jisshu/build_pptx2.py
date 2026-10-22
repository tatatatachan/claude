import math
import build_pptx as b   # slide 1 + helpers (b.sh is rebound to slide 2 below)
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt
from pptx.oxml.ns import qn
from lxml import etree
sl=b.prs.slides.add_slide(b.prs.slide_layouts[6]); b.sh=sh=sl.shapes
px,rect,text,fill,noline,shadow=b.px,b.rect,b.text,b.fill,b.noline,b.shadow
INK,SUB=b.INK,b.SUB
cx,cy,Ro,Ri=865,468,250,162

rect(0,0,1730,936,'E4F5E8').name='背景'
pts=[(0,0),(1730,0),(1730,132)]+[(1730*(1-i/240),132+10*i/240-20*math.sin(i/240*2*math.pi)) for i in range(241)]
fb=sh.build_freeform(px(pts[0][0]),px(pts[0][1])); fb.add_line_segments([(px(a),px(c)) for a,c in pts[1:]],close=True)
band=fb.convert_to_shape(); fill(band,'54C5AF'); noline(band); band.name='ヘッダー帯'
text(64,30,700,30,'施策　実習',20,'D9F6EE',True,spacing=4,name='小見出し')
text(64,60,1200,70,'どうブランドを育てていくか',46,'FFFFFF',True,name='タイトル')

steps=[ # a0,a1,color,dark,ring,head,desc,side,ly
 (270,360,'#3fb39d','1F7D6C','1 実習','共育実習を行う',['高校・専門学校・大学で','共育実習を実施します。'],'L',cy-122),
 (0,90,'#86d6c4','1F7D6C','2 声','声が集まる',['学生・先生・社員・','利用者さんの声が集まります。'],'R',cy-122),
 (90,180,'#f28c3c','C8651A','3 掲載','実習サイトに載せる',['プログラム内容と','口コミを掲載します。'],'R',cy+122),
 (180,270,'#2e9e8a','1F7D6C','4 行動','見て、動く',['学校は受け入れを判断し、','次の学生は参加を決めます。'],'L',cy+122),
]
c=sh.add_shape(MSO_SHAPE.OVAL,px(cx-136),px(cy-136),px(272),px(272)); fill(c,'FFFFFF'); noline(c); shadow(c); c.name='中央円'
text(cx-120,cy-70,240,32,'実習ブランドが',25,'2E9E8A',True,PP_ALIGN.CENTER,spacing=1,name='中央ミニ見出し')
text(cx-120,cy-38,240,66,'育つ循環',52,INK,True,PP_ALIGN.CENTER,name='中央タイトル')
p=rect(cx-90,cy+40,180,38,'F28C3C',MSO_SHAPE.ROUNDED_RECTANGLE,0.5); p.name='帯'
text(cx-90,cy+40,180,38,'1に戻って循環',19,'FFFFFF',True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,name='帯文字')

GAP=3; thick=(Ro-Ri)/(2*Ro)*100000
CW,CH=500,160
for (a0,a1,col,dark,ring,head,desc,side,ly) in steps:
    s_=sh.add_shape(MSO_SHAPE.BLOCK_ARC,px(cx-Ro),px(cy-Ro),px(2*Ro),px(2*Ro)); fill(s_,col); noline(s_); s_.name=f'環_{head}'
    av=s_._element.spPr.find(qn('a:prstGeom')).find(qn('a:avLst'))
    for g in list(av): av.remove(g)
    for n,v in (('adj1',((a0+GAP-90)%360)*60000),('adj2',((a1-GAP-90)%360)*60000),('adj3',thick)):
        etree.SubElement(av,qn('a:gd'),name=n,fmla=f'val {int(v)}')
    mid=(a0+a1)/2; rm=(Ro+Ri)/2
    mx=cx+rm*math.sin(math.radians(mid)); my=cy-rm*math.cos(math.radians(mid))
    rot=mid if (mid<=90 or mid>=270) else mid-180
    text(mx-80,my-20,160,40,ring,28,'FFFFFF',True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,spacing=2,rot=rot,name=f'環ラベル_{head}')
    off=ly-cy; R=Ro+14; dx=math.sqrt(R*R-off*off)
    left=side=='L'; x=64 if left else 1730-64-CW; y=ly-CH/2
    ex=cx-dx if left else cx+dx; cedge=x+CW if left else x
    ln=sh.add_connector(MSO_CONNECTOR.STRAIGHT,px(cedge),px(ly),px(ex),px(ly)); ln.line.color.rgb=b.rgb(col); ln.line.width=Pt(1.8); ln.name=f'接続線_{head}'
    d=sh.add_shape(MSO_SHAPE.OVAL,px(ex-6),px(ly-6),px(12),px(12)); fill(d,col); noline(d)
    card=rect(x,y,CW,CH,'FFFFFF',MSO_SHAPE.ROUNDED_RECTANGLE,0.14); shadow(card); card.name=f'カード_{head}'
    rect(x,y+30,8,CH-60,col,MSO_SHAPE.ROUNDED_RECTANGLE,0.5).name='カード色バー'
    text(x+36,y+20,440,30,'STEP '+ring[0],18,b.SUB,True,spacing=2,name='STEP')
    text(x+36,y+46,440,44,head,31,dark,True,name=f'要点_{head}')
    text(x+36,y+96,440,54,desc,20,INK,ls=1.15,name=f'説明_{head}')
rm=(Ro+Ri)/2
for a in (0,90,180,270):
    gx=cx+rm*math.sin(math.radians(a)); gy=cy-rm*math.cos(math.radians(a))
    o=sh.add_shape(MSO_SHAPE.OVAL,px(gx-14),px(gy-14),px(28),px(28)); fill(o,'FFFFFF'); noline(o); o.name='循環マーク白'
    t=math.radians(a); tx_,ty_=math.cos(t),math.sin(t); nx_,ny_=-ty_,tx_
    tip=(gx+8*tx_,gy+8*ty_); b1=(gx-4*tx_+4*nx_,gy-4*ty_+4*ny_); b2=(gx-4*tx_-4*nx_,gy-4*ty_-4*ny_)
    fb=sh.build_freeform(px(tip[0]),px(tip[1])); fb.add_line_segments([(px(b1[0]),px(b1[1])),(px(b2[0]),px(b2[1]))],close=True)
    ch=fb.convert_to_shape(); fill(ch,'2E9E8A'); noline(ch); ch.name='循環マーク矢印'

# bottom boxes
def box(x,w,title,lines,tc,bar,size,tcol):
    bx=rect(x,735,w,150,'FFFFFF',MSO_SHAPE.ROUNDED_RECTANGLE,0.1); shadow(bx); bx.name=f'ボックス_{title}'
    rect(x,761,8,98,bar,MSO_SHAPE.ROUNDED_RECTANGLE,0.5)
    text(x+34,754,w-60,34,title,26,tc,True,name=f'{title}_見出し')
    text(x+34,796,w-60,80,lines,size,tcol,ls=1.2,name=f'{title}_本文')
box(64,820,'サイトに載せるもの',['プログラム内容と口コミ。学校の先生が受け入れを判断でき、','学生が参加を決められる状態にする。'],INK,'F28C3C',21,INK)
box(910,756,'ほかの施策',['・関連大学・専門学校を増やす','・対応地域','・産官学で連携（次世代の育成を地域のみなさんで実施）'],SUB,'B8C4C0',19,SUB)

for el in sh._spTree.iter(qn('p:sp')):
    sp=el.find(qn('p:spPr'))
    if el.find(qn('p:style')) is not None and sp.find(qn('a:effectLst')) is None: etree.SubElement(sp,qn('a:effectLst'))
for er in sh._spTree.iter(qn('a:effectRef')): er.set('idx','0')
sl.notes_slide.notes_text_frame.text='実習ブランドが育つ循環：共育実習を行う→声が集まる→実習サイトに載せる→見て、動く→次の共育実習へ。'
b.prs.save('共育実習_図解.pptx')
