import math
import build_pptx2 as m2          # slides 1-2 (also loads helpers)
b=m2.b
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.util import Pt
from pptx.oxml.ns import qn
from lxml import etree
INK,SUB='26323A','5B6B73'
TEAL,TEAL2,TEALD,ORG,ORGD='54C5AF','3FB39D','1F7D6C','F28C3C','C8651A'

def new_slide(tag,title,notes=''):
    sl=b.prs.slides.add_slide(b.prs.slide_layouts[6]); b.sh=sl.shapes
    b.rect(0,0,1730,936,'E4F5E8').name='背景'
    pts=[(0,0),(1730,0),(1730,132)]+[(1730*(1-i/240),132+10*i/240-20*math.sin(i/240*2*math.pi)) for i in range(241)]
    fb=b.sh.build_freeform(b.px(pts[0][0]),b.px(pts[0][1])); fb.add_line_segments([(b.px(x),b.px(y)) for x,y in pts[1:]],close=True)
    band=fb.convert_to_shape(); b.fill(band,'54C5AF'); b.noline(band); band.name='ヘッダー帯'
    b.text(64,30,900,30,tag,20,'D9F6EE',True,spacing=4,name='小見出し')
    b.text(64,60,1500,70,title,46,'FFFFFF',True,name='タイトル')
    if notes: sl.notes_slide.notes_text_frame.text=notes
    return sl
def card(x,y,w,h,col=None,rad=0.12,fillc='FFFFFF',shadow=True):
    c=b.rect(x,y,w,h,fillc,MSO_SHAPE.ROUNDED_RECTANGLE,rad)
    if shadow: b.shadow(c)
    if col: b.rect(x,y+24,8,h-48,col,MSO_SHAPE.ROUNDED_RECTANGLE,0.5).name='色バー'
    return c
def note(t): b.text(64,894,1600,28,t,16,'7A8A87',name='注記')
def finish(sl):
    for el in sl.shapes._spTree.iter(qn('p:sp')):
        sp=el.find(qn('p:spPr'))
        if el.find(qn('p:style')) is not None and sp.find(qn('a:effectLst')) is None: etree.SubElement(sp,qn('a:effectLst'))
    for er in sl.shapes._spTree.iter(qn('a:effectRef')): er.set('idx','0')
def circle_num(x,y,d,n,col,fs=26):
    o=b.sh.add_shape(MSO_SHAPE.OVAL,b.px(x),b.px(y),b.px(d),b.px(d)); b.fill(o,col); b.noline(o); o.name='番号'
    b.text(x,y,d,d,str(n),fs,'FFFFFF',True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,name=f'番号文字{n}')

# ---------- 3: 全体像 ----------
sl=new_slide('実習応募サイト','実習応募サイトの全体像','エフビー介護サービス 実習応募サイト（制作中）の概要。')
p=b.rect(64,180,700,440,TEAL,MSO_SHAPE.ROUNDED_RECTANGLE,0.07); b.shadow(p); p.name='コンセプト'
b.text(104,206,620,30,'CARE PRACTICE',20,'D9F6EE',True,spacing=6,name='英字')
b.text(104,240,620,90,'FUN LIFE!',78,'FFFFFF',True,name='コピー1')
b.text(104,326,620,90,'FUN LOCAL!',78,'FFE08A',True,name='コピー2')
b.text(104,432,620,40,'生きがい、持ち寄る。地域の未来。',30,'FFFFFF',True,name='サブコピー')
b.text(104,494,620,110,['介護や老後の話は、なんだか暗い。','そんなイメージを、一緒にひっくり返したい。','日本一長寿の県・長野から、最高に楽しい仕事を。'],21,'FFFFFF',ls=1.25,name='メッセージ')
facts=[('目的',['介護の現場を実際に体験し、','仕事のイメージをつかむ'],None),('対象',['学生のみなさん','保護者・先生もご覧ください'],None),
       ('費用',['無償'],ORG),('期間・時間',['1DAY　9:30〜15:00ごろ','5DAYS　5日間じっくり'],None)]
for i,(lab,val,col) in enumerate(facts):
    x=800+(i%2)*434; y=180+(i//2)*226
    card(x,y,410,202,ORG if col else TEAL2)
    b.text(x+36,y+26,340,28,lab,20,SUB,True,spacing=2,name=f'ラベル_{lab}')
    if col: 
        b.text(x+36,y+62,340,80,'無償',64,ORGD,True,name='値_費用'); b.text(x+36,y+146,340,30,'実習は無料で参加できます',20,INK,name='値_費用補足')
    else: b.text(x+36,y+66,350,110,val,24 if lab=='目的' else 23,INK,True,ls=1.25,name=f'値_{lab}')
stats=[('109','拠点','信越・北関東に展開',TEALD),('7','種のサービス','ホーム・デイ・訪問など',TEALD),('2','つのプログラム','1DAY と 5DAYS',ORGD),('9','事業所','360°ツアーで見学できる',ORGD)]
for i,(n,u,s,c) in enumerate(stats):
    x=64+i*407; card(x,656,380,208)
    b.text(x+30,676,200,100,n,76,c,True,name=f'数字_{u}')
    b.text(x+30+(len(n)*46)+8,720,200,40,u,26,INK,True,name=f'単位_{u}')
    b.text(x+30,796,340,30,s,20,SUB,name=f'補足_{u}')
note('※ 対象・期間・時間などはサイト上でも仮・調整中の内容です。')
finish(sl)

# ---------- 4: エフビー流の学び ----------
sl=new_slide('エフビー流の学び','介護の現場を体験して学ぶ、4つのステップ','実習サイト「エフビー流の学び」の循環図を整理。')
pl=b.rect(565,176,600,64,TEAL,MSO_SHAPE.ROUNDED_RECTANGLE,0.5); pl.name='帯'
b.text(565,176,600,64,'介護の現場を体験して学ぶ',30,'FFFFFF',True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,name='帯文字')
steps=[('話して知る','利用者さんの人生にふれる',['昔の暮らしや仕事の話を聴く','笑顔が生まれる瞬間を見る','「ありがとう」をもらう手応え'],TEAL2,TEALD),
 ('見て学ぶ','職員の仕事を観察する',['声かけ・見守りの基本','介護職の一日の流れと仕事の幅','看護師など他職種との連携'],TEAL2,TEALD),
 ('やってみる','できることから体験する',['食事・レクのお手伝い','畑や行事など、地域にふれる','職員がそばでサポート'],ORG,ORGD),
 ('振り返る','気づきから、キャリアを考える',['感じたことを職員と話し合う','資格や仕事の道すじを知る','これからのキャリアを考える','自分に向いているかの見極め'],ORG,ORGD)]
for i,(h,s,bl,col,dk) in enumerate(steps):
    x=64+i*410; card(x,272,372,430)
    
    b.rect(x+30,276,312,6,col,MSO_SHAPE.ROUNDED_RECTANGLE,0.5).name='上端バー'
    circle_num(x+28,306,60,i+1,col,30)
    b.text(x+106,310,250,52,h,34,dk,True,anchor=MSO_ANCHOR.MIDDLE,name=f'見出し_{h}')
    b.text(x+30,392,320,56,s,21,SUB,True,ls=1.2,name=f'小見出し_{h}')
    ln=b.sh.add_connector(1,b.px(x+30),b.px(462),b.px(x+342),b.px(462)); ln.line.color.rgb=b.rgb('D7E6DF'); ln.line.width=Pt(1.5)
    b.text(x+30,486,320,300,['・'+t for t in bl],20,INK,ls=1.7,name=f'内容_{h}')
    if i<3:
        a=b.sh.add_shape(MSO_SHAPE.RIGHT_ARROW,b.px(x+372+4),b.px(460),b.px(30),b.px(44)); b.fill(a,'2E9E8A'); b.noline(a); a.name='矢印'
bk=b.sh.add_shape(MSO_SHAPE.LEFT_ARROW,b.px(64),b.px(750),b.px(1602),b.px(64)); b.fill(bk,'BFE6DC'); b.noline(bk); bk.adjustments[0]=0.5; bk.adjustments[1]=0.9; bk.name='戻り矢印'
b.text(220,750,1330,64,'④ 振り返りから、次の ① へ　体験をくり返しながら、学びを深めていく',22,TEALD,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,name='戻り矢印文字')
note('※ 文章はサイト上でも仮の内容です（決まり次第、差し替え）。')
finish(sl)

# ---------- 5: プログラム ----------
sl=new_slide('実習プログラム','1日で、5日間で。目的に合わせて選べる2つのプログラム','1DAY（9:30〜15:00ごろ）と5DAYS（5日間）。内容は調整中。')
def colhead(x,w,tag,txt,sub,col):
    h=b.rect(x,172,w,76,col,MSO_SHAPE.ROUNDED_RECTANGLE,0.3); h.name=f'帯_{tag}'
    b.text(x+28,172,200,76,tag,38,'FFFFFF',True,anchor=MSO_ANCHOR.MIDDLE,name=f'{tag}')
    b.text(x+210,182,w-230,34,txt,25,'FFFFFF',True,name=f'{tag}_見出し'); b.text(x+210,216,w-230,28,sub,19,'FFFFFF',name=f'{tag}_補足')
colhead(64,810,'1DAY','1日で、現場をのぞいてみる','9:30〜15:00ごろ　はじめての方に',TEAL2)
colhead(904,762,'5DAYS','5日間で、じっくり向き合う','DAY1〜DAY5　じっくり学びたい方に',ORG)
d1=[('9:30','受付とオープニング','メンターが温かく出迎え。会社紹介のショート動画で雰囲気を共有'),
 ('10:00','現場シャドーイング＆伴走体験','介護現場のリアルを体験。利用者様との「対話・人生のお話の傾聴」が中心'),
 ('12:00','ランチセッション','美味しいご飯と、先輩社員の地方暮らし・働き方のホンネトーク'),
 ('13:00','対話ワーク『100歳までの履歴書』','「どんな大人になっていたいか」から逆算し、大切にしたい「CARE」を言語化'),
 ('14:30','リフレクション＆フィードバック','メンターからの応援メッセージとフィードバックシートをお渡し')]
for i,(t,h,d) in enumerate(d1):
    y=266+i*120; card(64,y,810,108,TEAL2,0.14)
    chip=b.rect(92,y+30,120,48,'E3F4F0',MSO_SHAPE.ROUNDED_RECTANGLE,0.5); chip.name='時刻'
    b.text(92,y+30,120,48,t,24,TEALD,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,name=f'時刻_{t}')
    b.text(236,y+12,620,34,h,25,INK,True,name=f'要点_{h}'); b.text(236,y+50,620,52,d,18,SUB,ls=1.2,name=f'説明_{h}')
d5=[('DAY1','インストール','オリエンテーション・動画鑑賞・メンターと「実習中の目標」設定。利用者様の人生のお話に耳を傾ける'),
 ('DAY2-3','現場ケアと人間関係の構築','現場業務のサポートと利用者様との対話。メンターと毎日の振り返り'),
 ('DAY4','笑顔を届けるプチ企画','学生が考えた「利用者様を喜ばせる小さな企画」（お茶会・手書きメッセージ等）を実践'),
 ('DAY5','履歴書の発表会＆結びのランチ','「100歳までの履歴書」の発表、メンター全員からのメッセージ、修了式')]
for i,(t,h,d) in enumerate(d5):
    y=266+i*120; card(904,y,762,108,ORG,0.14)
    chip=b.rect(932,y+30,128,48,'FDEBDC',MSO_SHAPE.ROUNDED_RECTANGLE,0.5); chip.name='日'
    b.text(932,y+30,128,48,t,22,ORGD,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,name=f'日_{t}')
    b.text(1080,y+12,560,34,h,24,INK,True,name=f'要点_{h}'); b.text(1080,y+50,560,52,d,17,SUB,ls=1.2,name=f'説明_{h}')
bx=b.rect(904,746,762,118,'FFF3E4',MSO_SHAPE.ROUNDED_RECTANGLE,0.14); bx.name='掲載ボックス'
b.text(936,762,700,34,'実習の体験は、実習サイトへ',24,ORGD,True,name='掲載_見出し')
b.text(936,802,700,56,['実習の体験・感想を後日サイトに掲載。','Web掲載用の声としてストックします。'],19,INK,ls=1.25,name='掲載_本文')
note('※ プログラムの内容・日程はサイト上でも「調整中」です。')
finish(sl)

# ---------- 6: 流れ & サイトの中身 ----------
sl=new_slide('実習の流れ','はじめてでも安心。4つのステップと、サイトで届ける内容','申込から振り返りまでの4ステップと、サイトのコンテンツ構成。')
fl=[('申し込む','フォームから、行きたい日と事業所を選んで申込。質問だけでも大丈夫です。'),('事前の連絡','持ち物や服装、当日の集合場所をお知らせします。'),
    ('実習の当日','職員がそばでサポート。利用者さんとおしゃべりしながら、現場を体験。'),('振り返り','感じたことや気づいたことを、職員と話し合います。')]
for i,(h,d) in enumerate(fl):
    x=64+i*410; card(x,172,372,240)
    hb=b.rect(x,172,372,78,TEAL2 if i<3 else ORG,MSO_SHAPE.ROUNDED_RECTANGLE,0.25); 
    b.rect(x,210,372,40,TEAL2 if i<3 else ORG)  # square off bottom of header
    b.text(x+26,178,130,66,f'STEP 0{i+1}',22,'FFFFFF',True,anchor=MSO_ANCHOR.MIDDLE,spacing=1,name=f'STEP{i+1}')
    b.text(x+150,178,210,66,h,30,'FFFFFF',True,anchor=MSO_ANCHOR.MIDDLE,name=f'見出し_{h}')
    b.text(x+28,274,316,130,d,21,INK,ls=1.35,name=f'説明_{h}')
    if i<3:
        a=b.sh.add_shape(MSO_SHAPE.RIGHT_ARROW,b.px(x+376),b.px(268),b.px(30),b.px(44)); b.fill(a,'2E9E8A'); b.noline(a); a.name='矢印'
b.text(64,438,900,40,'実習サイトで届ける内容',30,INK,True,name='見出し_サイト')
tiles=[(['実習レポート'],'実習の1日の流れをたどる記事',TEAL2),(['職員の声'],'仕事のやりがいと、実習生へのメッセージ',TEAL2),(['人生の先輩の','メッセージ'],'ご利用者さんから、実習に来るみなさんへのひとこと',ORG),
       (['お知らせ'],'実習受け入れ拠点の追加など',TEAL2),(['事業所の','360°ツアー'],'事業所の中の様子をのぞける',TEAL2),(['経営者','インタビュー'],'社長に聞いた、これからのまちのカタチ（visions）',ORG)]
for i,(h,d,c) in enumerate(tiles):
    x=64+i*270; card(x,490,250,232,c,0.1)
    b.text(x+26,512,206,74,h,24,INK,True,ls=1.15,name=f'タイル_{h[0]}'); b.text(x+26,598,206,112,d,18,SUB,ls=1.3,name=f'タイル説明_{h[0]}')
st=b.rect(64,752,1602,112,'FFFFFF',MSO_SHAPE.ROUNDED_RECTANGLE,0.14); b.shadow(st); st.name='そのほか'
b.rect(64,776,8,64,TEAL2,MSO_SHAPE.ROUNDED_RECTANGLE,0.5)
b.text(100,768,200,40,'そのほか',24,TEALD,True,name='そのほか_見出し')
b.text(100,806,1540,50,'よくある質問　／　個人情報の取扱い（プライバシーポリシー）　／　気になるキーワードから記事を探す　／　実習の申込フォーム',20,INK,name='そのほか_本文')
note('※ 申込フォーム・SNS・問い合わせ先は準備中です。記事の内容はサンプルです。')
finish(sl)

# ---------- 7: 拠点 ----------
sl=new_slide('拠点','長野を中心に、信越・北関東に109の拠点','7種109拠点。長野57・新潟15・埼玉15・群馬12・栃木10。')
c=card(64,176,880,690)
b.text(100,196,700,40,'都県別の拠点数（計109拠点）',26,INK,True,name='グラフ見出し')
cd=CategoryChartData(); cats=['長野県','新潟県','埼玉県','群馬県','栃木県']; vals=[57,15,15,12,10]
cd.categories=cats; cd.add_series('拠点数',vals)
gf=b.sh.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,b.px(88),b.px(246),b.px(832),b.px(600),cd); ch=gf.chart; gf.name='拠点数グラフ'
ch.has_legend=False; ch.has_title=False
pl_=ch.plots[0]; pl_.gap_width=45; pl_.has_data_labels=True
dl=pl_.data_labels; dl.font.size=Pt(26); dl.font.bold=True; dl.font.name=b.FONT; dl.position=XL_LABEL_POSITION.OUTSIDE_END; dl.number_format='0'; dl.number_format_is_linked=False
for i in range(5):
    pt=pl_.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb=b.rgb(ORG if i==0 else '3FB39D')
va=ch.value_axis; va.visible=False; va.has_major_gridlines=False; va.maximum_scale=66; va.minimum_scale=0
ca=ch.category_axis; ca.tick_labels.font.size=Pt(22); ca.tick_labels.font.name=b.FONT; ca.format.line.color.rgb=b.rgb('C9D6D1'); ca.has_major_gridlines=False
c2=card(980,176,686,262)
b.text(1016,196,620,36,'提供しているサービス　7種',26,INK,True,name='サービス見出し')
svc=['有料老人ホーム','グループホーム','デイサービス','小規模多機能','看護小規模多機能型','訪問介護','訪問看護']
x0,y0=1016,252; 
pos=[(0,0),(1,0),(2,0),(0,1),(1,1),(0,2),(1,2)]
widths=[190,210,190]
for k,(t) in enumerate(svc):
    row=k//3 if k<6 else 2; col=k%3
    if k==6: col=0
    w=[200,200,200][col]; x=x0+col*212 ; y=y0+row*62
    ch_=b.rect(x,y,w,50,'E3F4F0',MSO_SHAPE.ROUNDED_RECTANGLE,0.5); ch_.name=f'サービス_{t}'
    b.text(x,y,w,50,t,18 if len(t)>8 else 20,TEALD,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,name=f'サービス文字_{t}')
c3=card(980,462,686,404,ORG,0.07)
b.text(1016,486,620,36,'360°ツアーで見学できる事業所　9拠点',24,INK,True,name='ツアー見出し')
rows=[('長野県','5','安曇野・松本・佐久・諏訪・上田'),('新潟県','1','糸魚川'),('栃木県','2','宇都宮・小山'),('埼玉県','1','羽生')]
for i,(pf,n,pl2) in enumerate(rows):
    y=552+i*78
    chip=b.rect(1016,y,120,48,'FDEBDC',MSO_SHAPE.ROUNDED_RECTANGLE,0.5); chip.name='県'
    b.text(1016,y,120,48,pf,20,ORGD,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,name=f'県_{pf}')
    b.text(1150,y,90,48,n+'拠点',22,INK,True,anchor=MSO_ANCHOR.MIDDLE,name=f'拠点数_{pf}')
    b.text(1250,y,390,48,pl2,20,SUB,anchor=MSO_ANCHOR.MIDDLE,name=f'所在地_{pf}')
note('※ 出典：エフビー介護サービス コーポレートサイト（360°ツアー）ほか、実習サイトの掲載内容。')
finish(sl)

b.prs.save('共育実習_図解.pptx')
print(len(b.prs.slides.__iter__.__self__._sldIdLst))
